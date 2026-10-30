"""Synthetic template for CI; the approved company's letterhead stays out of git."""
import json
from pathlib import Path
from docx import Document


def make_template(path):
    names=set(json.loads((Path(__file__).parents[1]/'config/iugn_contract_fields.json').read_text())['fields'])
    doc=Document()
    doc.add_paragraph("{%p if beneficiar_tip == 'PJ' %}")
    doc.add_paragraph('Beneficiar persoana juridica: {{ beneficiar_firma }}')
    doc.add_paragraph('Reprezentant {{ beneficiar_nume }}')
    doc.add_paragraph('{%p else %}')
    doc.add_paragraph('Beneficiar persoana fizica: {{ beneficiar_nume }}, CNP {{ beneficiar_cnp }}.')
    doc.add_paragraph('Act de identitate: CI {{ beneficiar_ci_nr }}')
    doc.add_paragraph('{%p endif %}')
    for name in sorted(names-{'beneficiar_firma','beneficiar_nume','beneficiar_ci_nr','beneficiar_cnp'}):
        doc.add_paragraph('{{ '+name+' }}')
    doc.add_paragraph('Anexa 4 Formular de apreciere demonstrativ')
    doc.save(path)
    return path
